import openpyxl

wb = openpyxl.load_workbook('DATA.xlsx', data_only=True)
sheet = wb['Agenda']
with open('agenda_dump.txt', 'w', encoding='utf-8') as f:
    f.write("Row 7:\n")
    for col in range(1, 15):
        f.write(f"{col}: {sheet.cell(row=7, column=col).value}\n")
    f.write("---\n")
    for row in range(10, 30):
        time_val = sheet.cell(row=row, column=2).value
        s1 = sheet.cell(row=row, column=5).value
        
        fill = sheet.cell(row=row, column=5).fill
        if fill and fill.start_color:
            c_type = fill.start_color.type
            if c_type == 'theme':
                s1_color = f"theme {fill.start_color.theme}"
            elif c_type == 'rgb':
                s1_color = f"rgb {fill.start_color.rgb}"
            else:
                s1_color = f"{c_type}"
        else:
            s1_color = "None"
        
        f.write(f"Row {row} | Time: {time_val} | S1: {s1} | S1 color: {s1_color}\n")
