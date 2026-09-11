import openpyxl
import codecs

wb = openpyxl.load_workbook('DATA.xlsx', data_only=True)
sheet = wb['Agenda']
with codecs.open('scratch/agenda_dump.txt', 'w', 'utf-8') as f:
    f.write("Row 7:\n")
    for col in range(1, 15):
        f.write(f"{col}: {sheet.cell(row=7, column=col).value}\n")
    f.write("---\n")
    for row in range(10, 30):
        time_val = sheet.cell(row=row, column=2).value
        s1 = sheet.cell(row=row, column=5).value
        s1_color = sheet.cell(row=row, column=5).fill.start_color.index if sheet.cell(row=row, column=5).fill else None
        f.write(f"Row {row} | Time: {time_val} | S1: {s1} | S1 color: {s1_color}\n")
