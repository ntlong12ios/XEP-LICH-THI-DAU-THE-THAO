import openpyxl
wb = openpyxl.load_workbook('DATA.xlsx', data_only=True)
sheet = wb['Agenda']
print("Row 7 columns:")
for col in range(1, 15):
    print(col, repr(sheet.cell(row=7, column=col).value))
print("---")
for row in range(10, 16):
    print("Row", row, "Time:", repr(sheet.cell(row=row, column=2).value), "Match S1:", repr(sheet.cell(row=row, column=5).value))
