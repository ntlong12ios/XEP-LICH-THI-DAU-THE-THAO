from fastapi import FastAPI, Depends, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy.orm import Session
from typing import List, Optional
import database
import models
import schemas

try:
    models.Base.metadata.create_all(bind=database.engine)
except Exception as e:
    print("Database connection error on startup:", e)

app = FastAPI(title="MT Tournament API")
from pydantic import BaseModel

import json

@app.post("/verify_admin")
def verify_admin(data: dict):
    if data.get("password") == "admin@1234":
        return {"success": True}
    return {"success": False}

@app.get("/led_state")
def get_led_state(db: Session = Depends(database.get_db)):
    state = db.query(models.AppState).filter(models.AppState.key == "led_state").first()
    if state and state.value:
        try:
            return json.loads(state.value)
        except:
            pass
    return {"type": "none", "data": None}

@app.post("/led_state")
def update_led_state(state: dict, db: Session = Depends(database.get_db)):
    state_str = json.dumps(state)
    db_state = db.query(models.AppState).filter(models.AppState.key == "led_state").first()
    if db_state:
        db_state.value = state_str
    else:
        db_state = models.AppState(key="led_state", value=state_str)
        db.add(db_state)
    db.commit()
    return {"success": True}

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€
# ROOT
# â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€
@app.get("/")
def read_root():
    return {"message": "Welcome to MT Tournament API"}

# â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€
# CATEGORIES
# â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€
@app.get("/categories", response_model=List[schemas.Category])
def get_categories(db: Session = Depends(database.get_db)):
    return db.query(models.Category).all()

# â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€
# TEAMS
# â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€
@app.get("/teams", response_model=List[schemas.Team])
def get_teams(db: Session = Depends(database.get_db)):
    return db.query(models.Team).all()

@app.get("/categories/{category_id}/teams", response_model=List[schemas.Team])
def get_teams_by_category(category_id: int, db: Session = Depends(database.get_db)):
    return db.query(models.Team).filter(models.Team.category_id == category_id).all()

import import_excel
from fastapi import File, UploadFile
import io

@app.post("/upload_excel")
async def upload_excel(file: UploadFile = File(...)):
    try:
        contents = await file.read()
        import_excel.run_import(file_obj=io.BytesIO(contents))
        return {"message": "Cập nhật dữ liệu từ file Excel thành công!"}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/reload_teams")
def reload_teams():
    try:
        import_excel.run_import()
        return {"message": "ÄÃ£ táº£i láº¡i thÃ´ng tin Ä‘á»™i tá»« file Excel!"}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Lá»—i: {str(e)}")

@app.put("/teams/{team_id}/draw_code", response_model=schemas.Team)
def update_team_draw_code(team_id: int, draw_code: str, db: Session = Depends(database.get_db)):
    # Remove this draw_code from any other team first (swap logic)
    existing = db.query(models.Team).filter(models.Team.draw_code == draw_code).first()
    
    team = db.query(models.Team).filter(models.Team.id == team_id).first()
    if not team:
        raise HTTPException(status_code=404, detail="Team not found")
    
    old_code = team.draw_code  # Save old code before overwriting
    
    if existing and existing.id != team_id:
        # Swap: give old code of this team to the other team
        existing.draw_code = old_code
        db.commit()

    team.draw_code = draw_code
    db.commit()
    db.refresh(team)
    return team

@app.delete("/teams/{team_id}/draw_code", response_model=schemas.Team)
def remove_team_draw_code(team_id: int, db: Session = Depends(database.get_db)):
    team = db.query(models.Team).filter(models.Team.id == team_id).first()
    if not team:
        raise HTTPException(status_code=404, detail="Team not found")
    team.draw_code = None
    db.commit()
    db.refresh(team)
    return team

# â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€
# BRACKET SLOTS (for draw screen)
# â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€
@app.get("/categories/{category_id}/bracket_slots", response_model=List[schemas.BracketSlot])
def get_bracket_slots(category_id: int, db: Session = Depends(database.get_db)):
    return db.query(models.BracketSlot).filter(
        models.BracketSlot.category_id == category_id
    ).order_by(models.BracketSlot.round_number, models.BracketSlot.slot_index).all()

@app.put("/bracket_slots/{slot_id}/assign", response_model=schemas.BracketSlot)
def assign_team_to_slot(slot_id: int, request: schemas.AssignTeamRequest, db: Session = Depends(database.get_db)):
    slot = db.query(models.BracketSlot).filter(models.BracketSlot.id == slot_id).first()
    if not slot:
        raise HTTPException(status_code=404, detail="Bracket slot not found")
    
    new_team_id = request.team_id
    
    # If there's already a team in this slot, unassign it
    if slot.team_id and slot.team_id != new_team_id:
        # Check if this team_id is in another slot in same category - if so swap
        other_slot = db.query(models.BracketSlot).filter(
            models.BracketSlot.category_id == slot.category_id,
            models.BracketSlot.team_id == new_team_id,
            models.BracketSlot.id != slot_id
        ).first()
        if other_slot:
            other_slot.team_id = slot.team_id
            db.commit()
    elif slot.team_id == new_team_id:
        # Already assigned, no change needed
        return slot
    else:
        # New assignment - check if team already in another slot
        other_slot = db.query(models.BracketSlot).filter(
            models.BracketSlot.category_id == slot.category_id,
            models.BracketSlot.team_id == new_team_id,
            models.BracketSlot.id != slot_id
        ).first()
        if other_slot:
            other_slot.team_id = None
            db.commit()
    
    slot.team_id = new_team_id
    db.commit()
    db.refresh(slot)
    return slot

@app.delete("/bracket_slots/{slot_id}/assign", response_model=schemas.BracketSlot)
def unassign_team_from_slot(slot_id: int, db: Session = Depends(database.get_db)):
    slot = db.query(models.BracketSlot).filter(models.BracketSlot.id == slot_id).first()
    if not slot:
        raise HTTPException(status_code=404, detail="Bracket slot not found")
    slot.team_id = None
    db.commit()
    db.refresh(slot)
    return slot

@app.post("/categories/{category_id}/generate_bracket", response_model=List[schemas.BracketSlot])
def generate_bracket(category_id: int, request: schemas.GenerateBracketRequest = None, db: Session = Depends(database.get_db)):
    """Generate bracket slots based on teams in this category"""
    if request is None:
        request = schemas.GenerateBracketRequest(format="knockout")
        
    category = db.query(models.Category).filter(models.Category.id == category_id).first()
    if not category:
        raise HTTPException(status_code=404, detail="Category not found")
    
    teams = db.query(models.Team).filter(models.Team.category_id == category_id).all()
    num_teams = len(teams)
    
    # Delete existing slots
    db.query(models.BracketSlot).filter(models.BracketSlot.category_id == category_id).delete()
    db.commit()
    
    slots = []
    
    if request.format == "round_robin":
        num_groups = request.num_groups or 8
        import string
        group_names = list(string.ascii_uppercase) # A, B, C...
        
        # Calculate teams per group (rough ceiling)
        teams_per_group = (num_teams + num_groups - 1) // num_groups
        if teams_per_group < 2: teams_per_group = 2
        
        slot_idx = 0
        for i in range(num_groups):
            g_name = group_names[i] if i < len(group_names) else f"B\u1ea3ng {i+1}"
            for pos in range(1, teams_per_group + 1):
                slot = models.BracketSlot(
                    category_id=category_id,
                    slot_code=f"B\u1ea3ng {g_name} - \u0110\u1ed9i {pos}",
                    round_number=1,
                    match_number=i+1, # match_number Ä‘áº¡i diá»‡n cho Bảng
                    slot_index=slot_idx,
                    position_in_match=pos,
                    label=f"B\u1ea3ng {g_name}"
                )
                db.add(slot)
                slots.append(slot)
                slot_idx += 1
                
    else:
        # Knockout
        slot_idx = 1
        for i in range(0, num_teams, 2):
            prefix = "BĐ Nam " if category.sport == "Bóng đá" else "Trận R1-"
            code = f"{prefix}{slot_idx:02d}" if category.sport != "Bóng đá" else f"{prefix}{slot_idx}"
            slot1 = models.BracketSlot(
                category_id=category_id,
                slot_code=code,
                round_number=1,
                match_number=slot_idx,
                slot_index=i,
                position_in_match=1,
                label=code
            )
            slot2 = models.BracketSlot(
                category_id=category_id,
                slot_code=code,
                round_number=1,
                match_number=slot_idx,
                slot_index=i+1,
                position_in_match=2,
                label=code
            )
            db.add(slot1)
            db.add(slot2)
            slots.extend([slot1, slot2])
            slot_idx += 1
    
    db.commit()
    for slot in slots:
        db.refresh(slot)
    
    return slots

# â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€
# MATCHES
# â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€
@app.get("/categories/{category_id}/matches", response_model=List[schemas.Match])
def get_matches(category_id: int, db: Session = Depends(database.get_db)):
    return db.query(models.Match).filter(
        models.Match.category_id == category_id
    ).order_by(models.Match.id).all()

@app.get("/matches", response_model=List[schemas.Match])
def get_all_matches(db: Session = Depends(database.get_db)):
    return db.query(models.Match).all()

@app.put("/matches/{match_id}/score", response_model=schemas.Match)
def update_match_score(match_id: int, request: schemas.UpdateScoreRequest, db: Session = Depends(database.get_db)):
    match = db.query(models.Match).filter(models.Match.id == match_id).first()
    if not match:
        raise HTTPException(status_code=404, detail="Match not found")
    
    match.score1 = request.score1
    match.score2 = request.score2
    match.status = "completed"
    
    # Determine winner and loser
    if request.score1 is not None and request.score2 is not None:
        winner_id = match.team1_id if request.score1 > request.score2 else match.team2_id
        loser_id = match.team2_id if request.score1 > request.score2 else match.team1_id
        
        # Advance winner to next match if set
        if match.next_match_id and winner_id:
            next_match = db.query(models.Match).filter(models.Match.id == match.next_match_id).first()
            if next_match:
                if next_match.team1_id is None:
                    next_match.team1_id = winner_id
                elif next_match.team2_id is None and next_match.team1_id != winner_id:
                    next_match.team2_id = winner_id

        # Advance loser to loser bracket if set
        if match.loser_next_match_id and loser_id:
            l_next_match = db.query(models.Match).filter(models.Match.id == match.loser_next_match_id).first()
            if l_next_match:
                if l_next_match.team1_id is None:
                    l_next_match.team1_id = loser_id
                elif l_next_match.team2_id is None and l_next_match.team1_id != loser_id:
                    l_next_match.team2_id = loser_id
    
    db.commit()
    db.refresh(match)
    return match

@app.put("/matches/{match_id}/schedule", response_model=schemas.Match)
def update_match_schedule(match_id: int, request: schemas.UpdateScheduleRequest, db: Session = Depends(database.get_db)):
    match = db.query(models.Match).filter(models.Match.id == match_id).first()
    if not match:
        raise HTTPException(status_code=404, detail="Match not found")
    
    match.scheduled_time = request.scheduled_time
    match.court = request.court
    
    db.commit()
    db.refresh(match)
    return match

import openpyxl
import datetime
@app.post("/load_agenda")
def load_agenda(db: Session = Depends(database.get_db)):
    import os
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    file_path = os.path.join(base_dir, "DATA.xlsx")
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb['Agenda']
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Lá»—i Ä‘á»c file Excel: {str(e)}")
        
    updates_made = 0
    for row_idx in range(11, sheet.max_row + 1):
        time_cell = sheet.cell(row=row_idx, column=2)
        if not time_cell.value: continue
            
        time_val = time_cell.value
        if isinstance(time_val, datetime.time) or isinstance(time_val, datetime.datetime):
            time_str = time_val.strftime("%H:%M")
        else:
            time_str = str(time_val).strip()
            if len(time_str) == 4 and ':' in time_str: time_str = '0' + time_str
                
        if ":" not in time_str: continue

        for col_idx in range(5, 14):
            cell = sheet.cell(row=row_idx, column=col_idx)
            match_code = cell.value
            if not match_code or not isinstance(match_code, str): continue
            match_code = str(match_code).strip()
            if not match_code: continue
                
            cat_id = None
            court = None
            
            if col_idx >= 12:
                cat_id = 1
                court = f"S\u00e2n b\u00f3ng {col_idx - 11}"
                if match_code == "Tranh 3-4" or match_code == "3-4":
                    match_code = "Tranh h\u1ea1ng 3"
            else:
                court = f"S\u00e2n {col_idx - 4}"
                if "N.N\u1eef" in match_code:
                    cat_id = 3
                elif "Nam TK" in match_code or "Nam BK" in match_code or "Nam CK" in match_code or "Nam 3-4" in match_code:
                    cat_id = 2
                elif "N\u1eef TK" in match_code or "N\u1eef BK" in match_code or "N\u1eef CK" in match_code or "N\u1eef 3-4" in match_code:
                    cat_id = 4
                else:
                    fill = cell.fill
                    if fill and fill.start_color:
                        c_type = fill.start_color.type
                        if c_type == 'theme':
                            if fill.start_color.theme == 3: cat_id = 2
                            elif fill.start_color.theme == 9 or fill.start_color.theme == 5: cat_id = 4
                        elif c_type == 'rgb':
                            if fill.start_color.rgb == 'FFFFFF00': cat_id = 3
                            
            if cat_id is not None:
                if match_code == "3-4" or match_code == "Tranh 3-4" or match_code == "Ba-T\u01b0" or match_code == "Tranh h\u1ea1ng 3":
                    if cat_id == 1:
                        match_code = "Tranh h\u1ea1ng 3"
                    elif cat_id == 2:
                        match_code = "Nam 3-4"
                    elif cat_id == 4:
                        match_code = "N\u1eef 3-4"
                    elif cat_id == 3:
                        match_code = "N.N\u1eef 3-4"

                from sqlalchemy import or_
                reversed_code = match_code
                if "-" in match_code:
                    parts = [p.strip() for p in match_code.split("-")]
                    if len(parts) == 2:
                        reversed_code = f"{parts[1]}-{parts[0]}"

                matches = db.query(models.Match).filter(
                    models.Match.category_id == cat_id,
                    or_(
                        models.Match.match_code == match_code,
                        models.Match.match_code == reversed_code
                    )
                ).all()
                for m in matches:
                    m.scheduled_time = time_str
                    m.court = court
                    updates_made += 1
    db.commit()
    return {"message": f"\u0110\u00e3 \u0111\u1ed3ng b\u1ed9 {updates_made} l\u1ecbch thi \u0111\u1ea5u t\u1eeb Agenda Excel!"}

import bracket_generator

@app.post("/categories/{category_id}/generate_matches")
def generate_matches_from_slots(category_id: int, db: Session = Depends(database.get_db)):
    """Generate matches from bracket slots after draw is complete using full tournament tree logic"""
    category = db.query(models.Category).filter(models.Category.id == category_id).first()
    if not category:
        raise HTTPException(status_code=404, detail="Category not found")
        
    slots = db.query(models.BracketSlot).filter(
        models.BracketSlot.category_id == category_id
    ).all()
    
    if not slots:
        raise HTTPException(status_code=400, detail="ChÆ°a cÃ³ dá»¯ liá»‡u bá»‘c thÄƒm! HÃ£y tiáº¿n hÃ nh bá»‘c thÄƒm trÆ°á»›c.")
    
    try:
        # Sá»­ dá»¥ng module chuyÃªn dá»¥ng Ä‘á»ƒ táº¡o sÆ¡ Ä‘á»“ giáº£i Ä‘áº¥u kÃ©p
        matches_created = bracket_generator.generate_matches(category, slots, db)
        return {"status": "ok", "matches_created": matches_created}
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Lá»—i táº¡o sÆ¡ Ä‘á»“: {str(e)}")



import os
import sys
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse

if getattr(sys, 'frozen', False):
    base_dir = sys._MEIPASS
else:
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

frontend_dist = os.path.join(base_dir, "frontend_dist")

# Mount assets directory
if os.path.exists(os.path.join(frontend_dist, "assets")):
    app.mount("/assets", StaticFiles(directory=os.path.join(frontend_dist, "assets")), name="assets")

@app.get("/{full_path:path}")
def serve_spa(full_path: str):
    path = os.path.join(frontend_dist, full_path)
    if full_path and os.path.isfile(path):
        return FileResponse(path)
    return FileResponse(os.path.join(frontend_dist, "index.html"))


# Vercel ASGI Prefix Stripper
class StripAPIPrefixMiddleware:
    def __init__(self, app):
        self.app = app
    async def __call__(self, scope, receive, send):
        if scope["type"] == "http" and scope["path"].startswith("/api/"):
            scope = dict(scope)
            scope["path"] = scope["path"][4:]
        await self.app(scope, receive, send)

app = StripAPIPrefixMiddleware(app)

