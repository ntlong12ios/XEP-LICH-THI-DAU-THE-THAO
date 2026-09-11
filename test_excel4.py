import openpyxl

wb = openpyxl.load_workbook('DATA.xlsx', data_only=True)
sheet = wb['Agenda']

matches = []
for row in range(10, 60):
    for c in range(5, 14):
        val = sheet.cell(row=row, column=c).value
        if val and isinstance(val, str) and "R1" in val:
            fill = sheet.cell(row=row, column=c).fill
            color = fill.start_color.theme if (fill and fill.start_color and fill.start_color.type == 'theme') else (fill.start_color.rgb if fill and fill.start_color and fill.start_color.type == 'rgb' else None)
            matches.append((val.encode('ascii', 'ignore').decode(), c, color))

print(matches)
